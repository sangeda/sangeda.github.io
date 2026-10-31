#!/usr/bin/env python3
"""Build the MUHAS Publications Observatory from OpenAlex.

The raw API snapshot is immutable. Derived tables add reporting-period fields
and are regenerated from raw data. Department/school assignment is deliberately
left for a versioned institutional crosswalk rather than inferred from names.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import html
import json
import shutil
import sqlite3
import os
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SITE = ROOT / "muhas-publications"
DOWNLOADS = SITE / "downloads"
DATA = SITE / "data"
RAW = ROOT / "muhas_publications_archive" / "raw"

INSTITUTION_NAME = "Muhimbili University of Health and Allied Sciences"
INSTITUTION_ALT_NAME = "Chuo Kikuu cha Afya na Sayansi Shirikishi Muhimbili"
OPENALEX_ID = "I154840374"
ROR = "027pr6c67"
REPOSITORY = "sangeda/sangeda.github.io"
PIPELINE_VERSION = "1.2"
START_DATE = date(1964, 1, 1)
USER_AGENT = "muhas-publications-observatory/1.2"
API_KEY_ENV = "OPENALEX_API_KEY"
# OpenAlex re-created most existing records on this date (index rebuild), so a record's
# created_date shows when it was first indexed only for quarters starting after it.
OPENALEX_CREATED_DATE_RESET = date(2025, 10, 10)
ADDED_DATE_NOTE = ("OpenAlex re-created most existing records on 10 October 2025, so \"added\" dates show "
                   "when a paper was indexed only for quarters starting after that date.")
# OpenAlex list results carry at most this many authorships per work; works at
# the cap are refetched from the single-work endpoint, which returns them all.
LIST_AUTHOR_CAP = 100
# Human-readable author strings show this many leading authors, then every
# MUHAS-affiliated author beyond them with their true position.
DISPLAY_LEADING_AUTHORS = 99

WORK_FIELDS = [
    "id", "doi", "ids", "title", "display_name", "publication_year", "publication_date",
    "type", "language", "cited_by_count", "fwci", "is_retracted", "is_paratext",
    "is_xpac", "countries_distinct_count", "institutions_distinct_count", "authorships",
    "is_authors_truncated", "primary_location", "open_access", "primary_topic", "topics",
    "keywords", "sustainable_development_goals", "corresponding_author_ids",
    "corresponding_institution_ids", "created_date", "updated_date",
]


def iso_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def short_id(value: str | None) -> str:
    return (value or "").rstrip("/").rsplit("/", 1)[-1]


def fiscal_period(d: date) -> tuple[str, str, date]:
    if 7 <= d.month <= 9:
        start, quarter, end = d.year, "Q1", date(d.year, 9, 30)
    elif 10 <= d.month <= 12:
        start, quarter, end = d.year, "Q2", date(d.year, 12, 31)
    elif 1 <= d.month <= 3:
        start, quarter, end = d.year - 1, "Q3", date(d.year, 3, 31)
    else:
        start, quarter, end = d.year - 1, "Q4", date(d.year, 6, 30)
    return f"FY{start}/{str(start + 1)[-2:]}", quarter, end


class OpenAlexAuthError(RuntimeError):
    """Raised on 401/403 so a missing or rejected API key fails the run at once."""


def with_api_key(url: str) -> str:
    key = os.environ.get(API_KEY_ENV, "").strip()
    if not key:
        return url
    return f"{url}{'&' if '?' in url else '?'}{urllib.parse.urlencode({'api_key': key})}"


def redact(url: str) -> str:
    return url.split("api_key=")[0].rstrip("?&") if "api_key=" in url else url


def api_json(url: str, attempts: int = 5) -> dict:
    request = urllib.request.Request(with_api_key(url), headers={"User-Agent": USER_AGENT})
    for attempt in range(attempts):
        try:
            with urllib.request.urlopen(request, timeout=90) as response:
                return json.load(response)
        except urllib.error.HTTPError as error:
            if error.code in (401, 403):
                hint = "" if os.environ.get(API_KEY_ENV) else f" ({API_KEY_ENV} is not set)"
                raise OpenAlexAuthError(f"OpenAlex refused the request with HTTP {error.code}{hint}: {redact(url)}") from None
            if attempt == attempts - 1:
                raise RuntimeError(f"OpenAlex HTTP {error.code} after {attempts} attempts: {redact(url)}") from None
            time.sleep(min(60, 5 * 2 ** attempt) if error.code == 429 else min(30, 2 ** attempt))
        except Exception as error:
            if attempt == attempts - 1:
                raise RuntimeError(f"OpenAlex request failed after {attempts} attempts ({type(error).__name__}): {redact(url)}") from None
            time.sleep(min(30, 2 ** attempt))
    raise RuntimeError("OpenAlex request failed")


def needs_full_authorships(work: dict) -> bool:
    return bool(work.get("is_authors_truncated")) or len(work.get("authorships") or []) >= LIST_AUTHOR_CAP


def complete_authorships(works: list[dict]) -> int:
    """Replace list records capped at 100 authorships with the full single-work record.

    Returns the number of works refetched. The list record's own truncation flag is
    kept as ``list_authors_truncated`` so the raw snapshot shows what was repaired.
    """
    refetched = 0
    select = urllib.parse.urlencode({"select": ",".join(WORK_FIELDS)})
    for index, work in enumerate(works):
        if not needs_full_authorships(work):
            continue
        full = api_json(f"https://api.openalex.org/works/{short_id(work.get('id'))}?{select}")
        if len(full.get("authorships") or []) < len(work.get("authorships") or []):
            raise RuntimeError(f"Single-work record for {short_id(work.get('id'))} has fewer authorships than the list record")
        full["list_authors_truncated"] = True
        works[index] = full
        refetched += 1
        time.sleep(0.12)
    return refetched


def fetch_openalex(as_of: date, raw_path: Path) -> tuple[list[dict], dict]:
    works: list[dict] = []
    cursor = "*"
    page = 0
    filter_value = (
        f"authorships.institutions.lineage:{OPENALEX_ID},"
        f"from_publication_date:{START_DATE.isoformat()},"
        f"to_publication_date:{as_of.isoformat()}"
    )
    while cursor:
        query = urllib.parse.urlencode({
            "filter": filter_value,
            "select": ",".join(WORK_FIELDS),
            "per-page": 200,
            "cursor": cursor,
        })
        payload = api_json(f"https://api.openalex.org/works?{query}")
        page += 1
        batch = payload.get("results", [])
        works.extend(batch)
        cursor = payload.get("meta", {}).get("next_cursor")
        if not batch:
            break
        print(f"OpenAlex page {page}: {len(works):,} works", flush=True)
        time.sleep(0.12)
    # Cursor paging over a live index can return a work twice; keep one record per id.
    unique = {short_id(w.get("id")): w for w in works}
    duplicates = len(works) - len(unique)
    works = list(unique.values())
    if duplicates:
        print(f"Dropped {duplicates:,} duplicate records returned while paging", flush=True)
    refetched = complete_authorships(works)
    print(f"Refetched {refetched:,} works capped at {LIST_AUTHOR_CAP} authorships", flush=True)
    raw_path.parent.mkdir(parents=True, exist_ok=True)
    with gzip.open(raw_path, "wt", encoding="utf-8") as stream:
        for work in works:
            stream.write(json.dumps(work, ensure_ascii=False) + "\n")
    return works, {
        "query_filter": filter_value,
        "openalex_reported_count": len(works),
        "pages_retrieved": page,
        "duplicate_records_dropped": duplicates,
        "works_refetched_for_full_authorships": refetched,
        "raw_path": str(raw_path.relative_to(ROOT)),
    }


def authors_display(names: list[str], muhas_flags: list[bool]) -> str:
    """All authors when short; otherwise the leading authors, then every MUHAS
    author beyond them with their true position, then the total count."""
    if len(names) <= LIST_AUTHOR_CAP:
        return " | ".join(names)
    parts = names[:DISPLAY_LEADING_AUTHORS]
    parts += [
        f"[MUHAS] {name} (pos. {position:,})"
        for position, (name, is_muhas) in enumerate(zip(names, muhas_flags), start=1)
        if position > DISPLAY_LEADING_AUTHORS and is_muhas
    ]
    return " | ".join(parts) + f" | … ({len(names):,} authors)"


FLAG_RULE_TEXT = (
    "Repository-hosted record (PubMed exempt) added to OpenAlex in 2026 or later "
    "with a publication date at least three years before it was added."
)


def is_backdated_deposit(work: dict) -> bool:
    """Flag (not exclude) repository deposits backdated relative to their OpenAlex creation."""
    source = (work.get("primary_location") or {}).get("source") or {}
    if source.get("type") != "repository" or "pubmed" in (source.get("display_name") or "").lower():
        return False
    created = (work.get("created_date") or "")[:4]
    year = work.get("publication_year")
    return created.isdigit() and int(created) >= 2026 and isinstance(year, int) and year <= int(created) - 3


def truthy(value) -> bool:
    return value is True or str(value).strip().lower() in ("true", "1", "yes")


def has_muhas_lineage(work: dict) -> bool:
    return any(
        OPENALEX_ID in (inst.get("lineage") or []) or short_id(inst.get("id")) == OPENALEX_ID
        for a in work.get("authorships") or [] for inst in a.get("institutions") or []
    )


def removed_status(ids: list[str]) -> dict:
    """Re-query works that left the set: deleted/merged, re-linked away from MUHAS, or still MUHAS."""
    found: dict[str, dict] = {}
    for i in range(0, len(ids), 50):
        query = urllib.parse.urlencode({
            "filter": "openalex:" + "|".join(ids[i:i + 50]),
            "select": "id,authorships",
            "per-page": 50,
        })
        for work in api_json(f"https://api.openalex.org/works?{query}").get("results", []):
            found[short_id(work.get("id"))] = work
        time.sleep(0.12)
    status = Counter()
    for wid in ids:
        work = found.get(wid)
        status["not_returned_deleted_or_merged" if work is None
               else "no_muhas_affiliation_now" if not has_muhas_lineage(work)
               else "still_muhas_other_reason"] += 1
    return dict(status)


def compare_works(previous: dict[str, dict], current: list[dict], query_api: bool = True) -> dict:
    """Aggregate differences between two raw snapshots (no titles)."""
    now = {short_id(w.get("id")): w for w in current}
    removed = sorted(set(previous) - set(now))
    added = sorted(set(now) - set(previous))
    report = {
        "previous_works": len(previous),
        "current_works": len(now),
        "added": len(added),
        "removed": len(removed),
        "removed_matching_backdated_rule": sum(is_backdated_deposit(previous[w]) for w in removed),
        "added_matching_backdated_rule": sum(is_backdated_deposit(now[w]) for w in added),
        "backdated_rule_previous": sum(map(is_backdated_deposit, previous.values())),
        "backdated_rule_current": sum(map(is_backdated_deposit, now.values())),
        "removed_by_source_type": dict(Counter(
            ((previous[w].get("primary_location") or {}).get("source") or {}).get("type") or "none" for w in removed
        ).most_common(5)),
    }
    if removed and query_api:
        report["removed_status_now"] = removed_status(removed)
    return report


def load_raw(path: Path) -> dict[str, dict]:
    with gzip.open(path, "rt", encoding="utf-8") as stream:
        return {short_id(w.get("id")): w for w in map(json.loads, stream)}


def join_names(items: list[dict] | None, field: str = "display_name") -> str:
    return " | ".join(str(item.get(field, "")).strip() for item in (items or []) if item.get(field))


def process_works(works: list[dict], retrieved_at: str, snapshot_id: str) -> tuple[list[dict], list[dict]]:
    work_rows: list[dict] = []
    author_rows: list[dict] = []
    for work in works:
        try:
            published = date.fromisoformat(work.get("publication_date") or f"{work['publication_year']}-01-01")
        except Exception:
            continue
        fy, quarter, period_end = fiscal_period(published)
        authorships = work.get("authorships") or []
        muhas_authors = []
        countries = set()
        corresponding_names = []
        corresponding_ids = set(work.get("corresponding_author_ids") or [])
        all_names: list[str] = []
        muhas_flags: list[bool] = []
        for sequence, authorship in enumerate(authorships, start=1):
            author = authorship.get("author") or {}
            institutions = authorship.get("institutions") or []
            has_muhas = any(OPENALEX_ID in (inst.get("lineage") or []) or short_id(inst.get("id")) == OPENALEX_ID for inst in institutions)
            name = author.get("display_name") or ""
            all_names.append(name or "(unnamed)")
            muhas_flags.append(has_muhas)
            if has_muhas and name:
                muhas_authors.append(name)
            if author.get("id") in corresponding_ids and name:
                corresponding_names.append(name)
            countries.update(authorship.get("countries") or [])
            author_rows.append({
                "snapshot_id": snapshot_id,
                "openalex_work_id": short_id(work.get("id")),
                "doi": (work.get("doi") or "").replace("https://doi.org/", ""),
                "publication_date": published.isoformat(),
                "fiscal_year": fy,
                "fiscal_quarter": quarter,
                "author_id": short_id(author.get("id")),
                "orcid": (author.get("orcid") or "").replace("https://orcid.org/", ""),
                "author_name": name,
                "author_position": authorship.get("author_position") or "",
                "author_sequence": sequence,
                "is_corresponding": bool(authorship.get("is_corresponding")),
                "has_muhas_affiliation": has_muhas,
                "institution_ids": " | ".join(short_id(x.get("id")) for x in institutions if x.get("id")),
                "institution_names": join_names(institutions),
                "institution_rors": " | ".join((x.get("ror") or "").replace("https://ror.org/", "") for x in institutions if x.get("ror")),
                "countries": " | ".join(sorted(authorship.get("countries") or [])),
                "school": "",
                "department": "",
                "curation_status": "unreviewed",
                "retrieved_at": retrieved_at,
            })
        primary_location = work.get("primary_location") or {}
        source = primary_location.get("source") or {}
        open_access = work.get("open_access") or {}
        primary_topic = work.get("primary_topic") or {}
        topics = work.get("topics") or []
        keywords = work.get("keywords") or []
        sdgs = work.get("sustainable_development_goals") or []
        corresponding_institutions = set(work.get("corresponding_institution_ids") or [])
        ids = work.get("ids") or {}
        work_rows.append({
            "snapshot_id": snapshot_id,
            "openalex_work_id": short_id(work.get("id")),
            "doi": (work.get("doi") or "").replace("https://doi.org/", ""),
            "pmid": short_id(ids.get("pmid")),
            "pmcid": short_id(ids.get("pmcid")),
            "title": work.get("title") or work.get("display_name") or "",
            "publication_date": published.isoformat(),
            "publication_year": published.year,
            "fiscal_year": fy,
            "fiscal_quarter": quarter,
            "period_end": period_end.isoformat(),
            "work_type": work.get("type") or "",
            "language": work.get("language") or "",
            "journal_or_source": source.get("display_name") or "",
            "source_openalex_id": short_id(source.get("id")),
            "issn_l": source.get("issn_l") or "",
            "publisher": source.get("host_organization_name") or "",
            "open_access_status": open_access.get("oa_status") or "",
            "is_open_access": bool(open_access.get("is_oa")),
            "open_access_url": open_access.get("oa_url") or "",
            "cited_by_count_at_snapshot": work.get("cited_by_count") or 0,
            "fwci_at_snapshot": work.get("fwci"),
            "is_retracted": bool(work.get("is_retracted")),
            "is_paratext": bool(work.get("is_paratext")),
            "is_xpac": bool(work.get("is_xpac")),
            "authors_count": len(authorships),
            "authors": authors_display(all_names, muhas_flags),
            "list_authors_truncated": bool(work.get("list_authors_truncated") or work.get("is_authors_truncated")),
            "muhas_authors_count": len(muhas_authors),
            "muhas_authors": " | ".join(muhas_authors),
            "corresponding_authors": " | ".join(corresponding_names),
            "has_muhas_corresponding_institution": any(short_id(x) == OPENALEX_ID for x in corresponding_institutions),
            "countries": " | ".join(sorted(countries)),
            "countries_count": work.get("countries_distinct_count") or len(countries),
            "institutions_count": work.get("institutions_distinct_count") or 0,
            "international_collaboration": len(countries) > 1,
            "primary_topic": primary_topic.get("display_name") or "",
            "primary_subfield": (primary_topic.get("subfield") or {}).get("display_name") or "",
            "primary_field": (primary_topic.get("field") or {}).get("display_name") or "",
            "primary_domain": (primary_topic.get("domain") or {}).get("display_name") or "",
            "topics": join_names(topics),
            "keywords": join_names(keywords),
            "sustainable_development_goals": join_names(sdgs),
            "flag_backdated_repository_deposit": is_backdated_deposit(work),
            "openalex_created_date": work.get("created_date") or "",
            "openalex_updated_date": work.get("updated_date") or "",
            "retrieved_at": retrieved_at,
            "school": "",
            "department": "",
            "curation_status": "unreviewed",
        })
    return work_rows, author_rows


def version_tuple(value) -> tuple[int, ...]:
    try:
        return tuple(int(x) for x in str(value).split("."))
    except ValueError:
        return (0,)


def parse_date(value) -> date | None:
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def fiscal_year_start(d: date) -> date:
    """MUHAS fiscal years run 1 July to 30 June."""
    return date(d.year if d.month >= 7 else d.year - 1, 7, 1)


def quarter_start(period_end: date) -> date:
    return date(period_end.year, period_end.month - 2, 1)


def previous_quarter_end(period_end: date) -> date:
    start = quarter_start(period_end)
    return date.fromordinal(start.toordinal() - 1)


def reporting_metrics(work_rows: list[dict], as_of: date) -> dict:
    """Fiscal-year and fiscal-quarter metrics for the page and the quarterly report.

    "Published" uses the publication date; "added" uses the date OpenAlex created
    the record, which shows works indexed late. Flagged deposits are counted in
    the totals and also reported excluding flags.
    """
    fy_now, q_now, end_now = fiscal_period(as_of)
    fy_start_now = fiscal_year_start(as_of)
    flagged = [truthy(r.get("flag_backdated_repository_deposit")) for r in work_rows]
    pub_key = Counter()
    pub_key_unflagged = Counter()
    added_key = Counter()
    added_key_unflagged = Counter()
    fiscal = Counter()
    fiscal_unflagged = Counter()
    for row, is_flagged in zip(work_rows, flagged):
        key = (row["fiscal_year"], row["fiscal_quarter"])
        pub_key[key] += 1
        fiscal[row["fiscal_year"]] += 1
        if not is_flagged:
            pub_key_unflagged[key] += 1
            fiscal_unflagged[row["fiscal_year"]] += 1
        created = parse_date(row.get("openalex_created_date"))
        if created:
            fy, q, _ = fiscal_period(created)
            added_key[(fy, q)] += 1
            if not is_flagged:
                added_key_unflagged[(fy, q)] += 1
    quarters = []
    end = end_now
    for _ in range(8):
        fy, q, _ = fiscal_period(end)
        reliable = quarter_start(end) > OPENALEX_CREATED_DATE_RESET
        quarters.append({
            "fiscal_year": fy, "fiscal_quarter": q,
            "period_start": quarter_start(end).isoformat(), "period_end": end.isoformat(),
            "partial": end > as_of,
            "published": pub_key[(fy, q)],
            "published_excluding_flagged": pub_key_unflagged[(fy, q)],
            "added": added_key[(fy, q)] if reliable else None,
            "added_excluding_flagged": added_key_unflagged[(fy, q)] if reliable else None,
        })
        end = previous_quarter_end(end)
    years = sorted(fiscal)
    fiscal_years = [
        {"fiscal_year": fy, "publications": fiscal[fy], "publications_excluding_flagged": fiscal_unflagged[fy],
         "partial": fy == fy_now and as_of < date(fy_start_now.year + 1, 6, 30)}
        for fy in years
    ]
    complete_fy = fy_now if as_of >= date(fy_start_now.year + 1, 6, 30) else fiscal_period(date(fy_start_now.year - 1, 7, 1))[0]
    topics = Counter(r["primary_topic"] for r in work_rows if r["fiscal_year"] == complete_fy and r.get("primary_topic"))
    latest = quarters[0]
    q_start = date.fromisoformat(latest["period_start"])
    added_earlier = sum(
        1 for r in work_rows
        if (c := parse_date(r.get("openalex_created_date"))) and q_start <= c <= end_now
        and (pd := parse_date(r.get("publication_date"))) and pd < q_start
    )
    return {
        "as_of_date": as_of.isoformat(),
        "current_fiscal_year": fy_now,
        "latest_quarter": {**latest, "added_published_earlier": added_earlier},
        "quarters": quarters,
        "fiscal_years": fiscal_years,
        "works_before_current_fiscal_year": sum(
            1 for r in work_rows if (pd := parse_date(r.get("publication_date"))) and pd < fy_start_now
        ),
        "topics_fiscal_year": complete_fy,
        "added_date_note": ADDED_DATE_NOTE,
        "top_topics": [{"topic": t, "publications": n} for t, n in topics.most_common(25)],
    }


QUARTER_COLUMNS = [
    "openalex_work_id", "doi", "pmid", "title", "publication_date", "date_is_1_january",
    "journal_or_source", "work_type", "open_access_status", "muhas_authors_count", "muhas_authors",
    "corresponding_authors", "has_muhas_corresponding_institution", "international_collaboration",
    "primary_topic", "primary_field", "flag_backdated_repository_deposit", "openalex_created_date",
]


def write_quarterly_report(path: Path, work_rows: list[dict], manifest: dict) -> dict:
    """The DRPI workbook: what was published, and what was newly indexed, in the latest quarter."""
    as_of = date.fromisoformat(manifest["as_of_date"])
    metrics = reporting_metrics(work_rows, as_of)
    latest = metrics["latest_quarter"]
    q_start, q_end = date.fromisoformat(latest["period_start"]), date.fromisoformat(latest["period_end"])

    def pick(row: dict) -> dict:
        out = {c: row.get(c, "") for c in QUARTER_COLUMNS}
        out["date_is_1_january"] = str(row.get("publication_date", ""))[5:10] == "01-01"
        return out

    published = sorted(
        (r for r in work_rows if (r["fiscal_year"], r["fiscal_quarter"]) == (latest["fiscal_year"], latest["fiscal_quarter"])),
        key=lambda r: (str(r["publication_date"]), str(r["title"]).lower()),
    )
    earlier = sorted(
        (r for r in work_rows
         if (c := parse_date(r.get("openalex_created_date"))) and q_start <= c <= q_end
         and (pd := parse_date(r.get("publication_date"))) and pd < q_start),
        key=lambda r: (str(r["publication_date"]), str(r["title"]).lower()),
        reverse=True,
    )
    label = f"{latest['fiscal_year']} {latest['fiscal_quarter']}" + (" (to date)" if latest["partial"] else "")
    tag = f"muhas-pubs-{manifest['as_of_date']}-{manifest['snapshot_status']}"
    counts = {
        "quarter": label,
        "published_in_quarter": len(published),
        "published_dated_1_january": sum(str(r["publication_date"])[5:10] == "01-01" for r in published),
        "published_flagged": sum(truthy(r.get("flag_backdated_repository_deposit")) for r in published),
        "added_in_quarter_published_earlier": len(earlier),
        "added_earlier_flagged": sum(truthy(r.get("flag_backdated_repository_deposit")) for r in earlier),
    }
    wb = Workbook()
    readme = wb.active
    readme.title = "Read_me"
    readme.append(["MUHAS Publications Observatory: quarterly report", ""])
    for item in [
        ("Quarter", label),
        ("Period", f"{q_start.isoformat()} to {q_end.isoformat()}"),
        ("Snapshot", manifest["snapshot_id"]),
        ("Release (cite this)", f"https://github.com/{REPOSITORY}/releases/tag/{tag}"),
        ("Snapshot status", manifest["snapshot_status"]),
        ("Works published in the quarter", counts["published_in_quarter"]),
        ("  of which dated 1 January", counts["published_dated_1_january"]),
        ("  of which flagged repository deposits", counts["published_flagged"]),
        ("Works added to OpenAlex in the quarter but published earlier", counts["added_in_quarter_published_earlier"]),
        ("  of which flagged repository deposits", counts["added_earlier_flagged"]),
        ("Definition: published", "Publication date within the MUHAS fiscal quarter (fiscal year 1 July to 30 June)."),
        ("Definition: added", "Date OpenAlex created the record. Papers are often indexed weeks or months after publication, "
                              "so later reports list late additions for earlier quarters."),
        ("Caveat: added dates", ADDED_DATE_NOTE),
        ("Caveat: 1 January dates", "Some records carry only a publication year, which OpenAlex dates 1 January; these fall in Q3 "
                                    "and are marked in the date_is_1_january column."),
        ("Caveat: flags", FLAG_RULE_TEXT + " Flagged works are under review and are not necessarily invalid."),
        ("Caveat: pilot", "Pilot dataset from OpenAlex (institution I154840374, ROR 027pr6c67), cleaned progressively. "
                          "Not an audited institutional return."),
        ("Schools and departments", "Not yet assigned; they will be added once the staff affiliation crosswalk is approved."),
    ]:
        readme.append(list(item))
    style_sheet(readme)
    readme.column_dimensions["A"].width = 58
    readme.column_dimensions["B"].width = 100
    for name, rows in (("Quarter_published", published), ("Added_published_earlier", earlier)):
        ws = wb.create_sheet(name)
        selected = [pick(r) for r in rows]
        if name == "Added_published_earlier":
            for out, r in zip(selected, rows):
                out["published_fiscal_quarter"] = f"{r['fiscal_year']} {r['fiscal_quarter']}"
        if selected:
            append_rows(ws, selected)
        else:
            ws.append(["No works in this category for the quarter."])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return counts


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def summaries(work_rows: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    annual = Counter(row["publication_year"] for row in work_rows)
    quarterly = Counter((row["fiscal_year"], row["fiscal_quarter"], row["period_end"]) for row in work_rows)
    types = Counter(row["work_type"] or "unknown" for row in work_rows)
    annual_rows = [{"publication_year": year, "publications": count} for year, count in sorted(annual.items())]
    quarter_rows = [
        {"fiscal_year": fy, "fiscal_quarter": q, "period_end": end, "publications": count}
        for (fy, q, end), count in sorted(quarterly.items())
    ]
    type_rows = [{"work_type": kind, "publications": count} for kind, count in types.most_common()]
    return annual_rows, quarter_rows, type_rows


def write_sqlite(path: Path, works: list[dict], authorships: list[dict], manifest: dict) -> None:
    if path.exists():
        path.unlink()
    journal_path = Path(str(path) + "-journal")
    if journal_path.exists():
        journal_path.unlink()
    connection = sqlite3.connect(path)
    try:
        # A self-contained downloadable database must not depend on a rollback
        # journal that can be separated from it during copying or download.
        connection.execute("PRAGMA journal_mode=OFF")
        connection.execute("PRAGMA synchronous=OFF")
        for table, rows in (("works", works), ("authorships", authorships)):
            columns = list(rows[0])
            connection.execute(f'CREATE TABLE {table} ({", ".join(f"[{c}] TEXT" for c in columns)})')
            connection.executemany(
                f'INSERT INTO {table} VALUES ({", ".join("?" for _ in columns)})',
                [[str(row.get(col, "")) if row.get(col) is not None else "" for col in columns] for row in rows],
            )
        connection.execute("CREATE TABLE snapshots (snapshot_id TEXT, metadata_json TEXT)")
        connection.execute("INSERT INTO snapshots VALUES (?, ?)", (manifest["snapshot_id"], json.dumps(manifest)))
        connection.execute("CREATE INDEX idx_works_date ON works(publication_date)")
        connection.execute("CREATE INDEX idx_auth_work ON authorships(openalex_work_id)")
        connection.execute("CREATE INDEX idx_auth_author ON authorships(author_id)")
        connection.commit()
        result = connection.execute("PRAGMA integrity_check").fetchone()[0]
        if result != "ok":
            raise RuntimeError(f"SQLite integrity check failed: {result}")
        if connection.execute("SELECT COUNT(*) FROM works").fetchone()[0] != len(works):
            raise RuntimeError("SQLite works row count does not match source rows")
    finally:
        connection.close()


def style_sheet(ws) -> None:
    fill = PatternFill("solid", fgColor="0F716C")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in range(1, min(ws.max_column, 40) + 1):
        values = [str(ws.cell(row=r, column=column).value or "") for r in range(1, min(ws.max_row, 200) + 1)]
        ws.column_dimensions[get_column_letter(column)].width = min(45, max(10, max(map(len, values), default=10) + 2))


def append_rows(ws, rows: list[dict]) -> None:
    if not rows:
        return
    ws.append(list(rows[0]))
    for row in rows:
        ws.append([row.get(column, "") for column in rows[0]])
    style_sheet(ws)


def write_excel(path: Path, works: list[dict], authorships: list[dict], annual: list[dict], quarters: list[dict], types: list[dict], manifest: dict) -> None:
    wb = Workbook(write_only=False)
    summary = wb.active
    summary.title = "Executive_Summary"
    summary.append(["MUHAS Publications Observatory", "Value"])
    summary_rows = [
        ("Institution", INSTITUTION_NAME), ("Alternate name", INSTITUTION_ALT_NAME),
        ("ROR", ROR), ("OpenAlex ID", OPENALEX_ID), ("Snapshot ID", manifest["snapshot_id"]),
        ("Retrieved at", manifest["retrieved_at"]), ("Period through", manifest["as_of_date"]),
        ("Works", len(works)), ("Authorship rows", len(authorships)),
        ("Important", "Pilot OpenAlex-derived dataset, cleaned progressively; school and department assignments await the approved affiliation crosswalk."),
    ]
    for row in summary_rows:
        summary.append(row)
    style_sheet(summary)
    for name, rows in (
        ("Publications", works), ("Authorships", authorships),
        ("Fiscal_Year_Summary", [{"fiscal_year": fy, "publications": n} for fy, n in sorted(Counter(r["fiscal_year"] for r in works).items())]),
        ("Annual_Summary", annual),
        ("Quarterly_Summary", quarters), ("Type_Summary", types),
        ("Snapshot_Metadata", [{"field": k, "value": json.dumps(v) if isinstance(v, (dict, list)) else v} for k, v in manifest.items()]),
    ):
        append_rows(wb.create_sheet(name), rows)
    curation = wb.create_sheet("Affiliation_Crosswalk")
    append_rows(curation, [{
        "author_id": "", "orcid": "", "author_name": "", "school": "", "department": "",
        "effective_from": "", "effective_to": "", "curation_status": "", "evidence": "",
        "curated_by": "", "curated_at": "",
    }])
    dictionary = wb.create_sheet("Data_Dictionary")
    append_rows(dictionary, [
        {"field": "fiscal_year / fiscal_quarter", "definition": "MUHAS fiscal year (1 July to 30 June) and quarter (Q1 Jul-Sep, Q2 Oct-Dec, Q3 Jan-Mar, Q4 Apr-Jun) of the publication date."},
        {"field": "period_end", "definition": "Official end date of the MUHAS fiscal quarter containing the publication date."},
        {"field": "Annual_Summary", "definition": "Calendar-year counts (January to December), kept for comparison with journal and Scopus figures."},
        {"field": "retrieved_at", "definition": "Actual UTC timestamp when the source API data were retrieved."},
        {"field": "snapshot_id", "definition": "Immutable identifier for the extraction run."},
        {"field": "curation_status", "definition": "Institutional review status; records default to unreviewed until curated."},
        {"field": "cited_by_count_at_snapshot", "definition": "OpenAlex citation count at retrieval; not publication productivity."},
        {"field": "school / department", "definition": "Blank until assigned through the versioned MUHAS affiliation crosswalk."},
        {"field": "authors", "definition": "All authors; for works with more than 100, the first 99, then every MUHAS author beyond them with true position, then the total."},
        {"field": "list_authors_truncated", "definition": "OpenAlex list results capped this work at 100 authorships; the full record was refetched, so authorship rows are complete."},
        {"field": "author_sequence", "definition": "1-based position of the author in the work's full author list."},
        {"field": "flag_backdated_repository_deposit", "definition": FLAG_RULE_TEXT + " Flagged for review, not excluded."},
        {"field": "pmid / pmcid", "definition": "PubMed and PubMed Central identifiers recorded by OpenAlex, where present."},
    ])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def changes_html(changes: dict | None) -> str:
    """Neutral note on how this snapshot differs from the previous final one."""
    if not changes:
        return ""
    tag = str(changes.get("previous_snapshot", ""))
    date_match = tag[11:21] if tag.startswith("muhas-pubs-") else tag
    status = changes.get("removed_status_now") or {}
    parts = [f"{changes.get('added', 0):,} works were added and {changes.get('removed', 0):,} are no longer included."]
    detail = []
    if status.get("not_returned_deleted_or_merged"):
        detail.append(f"{status['not_returned_deleted_or_merged']:,} are no longer returned by OpenAlex (deleted or merged into another record)")
    if status.get("no_muhas_affiliation_now"):
        detail.append(f"{status['no_muhas_affiliation_now']:,} are no longer linked to MUHAS")
    if changes.get("removed_matching_backdated_rule"):
        detail.append(f"{changes['removed_matching_backdated_rule']:,} had matched the repository-deposit flag")
    if detail:
        parts.append("Of those no longer included, " + "; ".join(detail) + ".")
    parts.append("Counts for earlier years can change between snapshots as OpenAlex corrects its records; "
                 "each dated snapshot is archived unchanged as a GitHub Release.")
    return (f'<p class="changes-note"><strong>Changes since the {html.escape(date_match)} snapshot.</strong> '
            f'{html.escape(" ".join(parts))}</p>')


def build_dashboard(work_rows: list[dict], annual: list[dict], quarter_rows: list[dict], type_rows: list[dict], manifest: dict) -> None:
    DATA.mkdir(parents=True, exist_ok=True)
    unflagged = [row for row in work_rows if not truthy(row.get("flag_backdated_repository_deposit"))]
    flagged = len(work_rows) - len(unflagged)
    annual_unflagged, quarters_unflagged, _ = summaries(unflagged)
    as_of = date.fromisoformat(manifest["as_of_date"])
    reporting = reporting_metrics(work_rows, as_of)
    latest = reporting["latest_quarter"]
    q_label = f"{latest['fiscal_year']} {latest['fiscal_quarter']}" + (" to date" if latest["partial"] else "")
    summary = {
        "reporting": reporting,
        "flagged_backdated_deposits": flagged,
        "flag_rule": FLAG_RULE_TEXT,
        "annual_excluding_flagged": annual_unflagged,
        "quarters_excluding_flagged": quarters_unflagged,
        "institution": INSTITUTION_NAME,
        "openalex_id": OPENALEX_ID,
        "ror": ROR,
        "snapshot": manifest,
        "total_works": len(work_rows),
        "historical_1964_2025": sum(row["publication_year"] <= 2025 for row in work_rows),
        "annual": annual,
        "quarters": quarter_rows,
        "types": type_rows[:15],
        "open_access": dict(Counter(row["open_access_status"] or "unknown" for row in work_rows)),
        "top_topics": [
            {"topic": topic, "publications": count}
            for topic, count in Counter(row["primary_topic"] for row in work_rows if row["primary_topic"]).most_common(25)
        ],
    }
    (DATA / "summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    flag_note = (
        f'<p class="flag-note"><label><input type="checkbox" id="exclude-flagged"> Exclude {flagged:,} flagged repository deposits</label> '
        f'<span>{html.escape(FLAG_RULE_TEXT)} These records are under review and are counted in the totals unless excluded; '
        f'flagging does not mean they are invalid.</span></p>'
    ) if flagged else ""
    changes_note = changes_html(manifest.get("changes_since_previous"))
    # Describe what the files on offer actually contain: snapshots before pipeline 1.1
    # lack PubMed IDs, flags and complete author lists.
    current_files = version_tuple(manifest.get("pipeline_version", "1.0")) >= (1, 1)
    # The quarterly report lives in the dated release, so the link survives the shared deploy.
    quarterly_card = (
        f'<a href="https://github.com/{REPOSITORY}/releases/download/'
        f'muhas-pubs-{html.escape(manifest["as_of_date"])}-{html.escape(manifest["snapshot_status"])}/'
        f'{html.escape(manifest["snapshot_id"])}_quarterly_report.xlsx"><strong>Quarterly report (Excel)</strong>'
        f'<span>Works published in {html.escape(q_label)}, and works added to OpenAlex in the quarter but published earlier</span></a>'
    ) if version_tuple(manifest.get("pipeline_version", "1.0")) >= (1, 2) else ""
    cards = (
        ("Publications, authorships, summaries, data dictionary and metadata",
         "One row per work, with PubMed IDs and review flags",
         "One row per author and work, complete author lists")
        if current_files else
        ("Publications, authorships, summaries, data dictionary and metadata",
         "One row per OpenAlex work",
         "One row per author–publication relationship")
    )
    release_tag = f"muhas-pubs-{manifest['as_of_date']}-{manifest['snapshot_status']}"
    release_link = (f'<a href="https://github.com/{REPOSITORY}/releases/tag/{html.escape(release_tag)}">'
                    f'<code>{html.escape(release_tag)}</code></a>')
    page = f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>MUHAS Publications Observatory</title><meta name="description" content="Reproducible OpenAlex-based MUHAS publication reporting by year and fiscal quarter.">
<link rel="stylesheet" href="../assets/style.css"><style>.flag-note{{display:flex;flex-wrap:wrap;gap:.4rem 1rem;align-items:baseline;margin:0 0 1.2rem;font-size:.95rem}}.flag-note label{{font-weight:700;white-space:nowrap}}.flag-note span{{opacity:.8}}.changes-note{{margin:0 0 1.2rem;font-size:.95rem;max-width:62rem}}.snapshot-line{{margin:1.2rem 0 0;font-size:.9rem;opacity:.75}}.table-note{{margin:.8rem 0 0;font-size:.85rem;opacity:.75}}.mini-bars.fy .mini-bar-row{{grid-template-columns:5.2rem 1fr 48px}}.mini-bars.fy .mini-bar-row span{{white-space:nowrap}}#muhas-summary td:first-child{{white-space:nowrap}}</style><link rel="icon" href="../assets/favicon.svg" type="image/svg+xml"></head><body>
<header class="site-header"><div class="nav-wrap"><a class="brand" href="../index.html"><span class="brand-mark">MU</span><span>MUHAS Publications Observatory</span></a><nav class="site-nav always"><a href="../index.html">Profile</a><a href="#downloads">Downloads</a><a href="#methods">Methods</a></nav></div></header>
<main><section class="pub-hero"><p class="eyebrow">Institutional research intelligence · Pilot</p><h1>MUHAS publications, <span>preserved and reportable.</span></h1>
<p class="hero-lead">A pilot observatory of MUHAS publications from OpenAlex, refreshed after each fiscal quarter. Records are cleaned progressively, so counts can change between snapshots; school and department reporting will follow once the staff affiliation crosswalk is approved.</p>
<div class="hero-stats"><div><strong>{len(work_rows):,}</strong><span>works through {manifest['as_of_date']}</span></div><div><strong>{latest['published']:,}</strong><span>published in {html.escape(q_label)}</span></div><div><strong>{"—" if latest['added'] is None else f"{latest['added']:,}"}</strong><span>added to OpenAlex in {html.escape(q_label)}</span></div><div><strong>{reporting['works_before_current_fiscal_year']:,}</strong><span>published before {html.escape(reporting['current_fiscal_year'])}</span></div></div>
<p class="snapshot-line">{html.escape(manifest['snapshot_status'].title())} snapshot · retrieved {html.escape(manifest['retrieved_at'][:10])} · fiscal year runs 1 July to 30 June</p></section>
<section class="section tinted"><div class="section-heading"><div><p class="section-kicker">Reporting</p><h2>Fiscal-year and quarterly views</h2></div><p>Reporting follows the MUHAS fiscal year (1 July to 30 June). "Published" uses the publication date; "added" is when OpenAlex indexed the record, which shows papers that arrive late.</p></div>{flag_note}{changes_note}<div id="muhas-summary" class="observatory-grid"></div></section>
<section id="downloads" class="section"><div class="section-heading"><div><p class="section-kicker">Downloads</p><h2>Data downloads</h2></div><p>Files for the snapshot shown on this page. The LATEST files are replaced every quarter; to cite a fixed version, use the dated release {release_link}, which also holds the raw OpenAlex records. Excel is formatted for review; CSV and SQLite suit R, Python and Power BI.</p></div>
<div class="download-grid">{quarterly_card}<a href="downloads/MUHAS_Publications_LATEST.xlsx"><strong>Excel workbook</strong><span>{cards[0]}</span></a><a href="downloads/MUHAS_Publications_LATEST.csv"><strong>Publications CSV</strong><span>{cards[1]}</span></a><a href="downloads/MUHAS_Authorships_LATEST.csv"><strong>Authorships CSV</strong><span>{cards[2]}</span></a><a href="downloads/MUHAS_Publications_LATEST.sqlite"><strong>SQLite database</strong><span>Works, authorships and snapshot metadata</span></a></div></section>
<section id="methods" class="section dark"><div class="section-heading"><div><p class="section-kicker">Methods and cautions</p><h2>A pilot, curated progressively.</h2></div></div><div class="prose"><p>The extraction uses OpenAlex institution lineage <code>{OPENALEX_ID}</code> and ROR <code>{ROR}</code>. Works with more than 100 authors are fetched in full, so MUHAS authors late in long author lists are counted. Repository deposits added to OpenAlex long after their stated publication date are flagged for review and can be excluded from the charts. School and department fields remain blank until MUHAS approves a time-aware author affiliation crosswalk.</p><p>This is a pilot. OpenAlex records can be incomplete, delayed or incorrectly affiliated, and both OpenAlex and this observatory correct them progressively. Counts are research-intelligence evidence, not an audited institutional return; each dated snapshot is archived unchanged for reproducibility.</p></div></section></main>
<footer><p>MUHAS Publications Observatory · Pilot, refreshed quarterly</p><p>Snapshot {html.escape(manifest['snapshot_id'])}</p></footer><script src="assets/observatory.js"></script></body></html>'''
    (SITE / "index.html").write_text(page, encoding="utf-8")


def build_js() -> None:
    assets = SITE / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    script = '''(function(){const box=()=>document.querySelector("#exclude-flagged");let data=null;
const esc=t=>String(t).replace(/[&<>"]/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
function render(){const d=data;const ex=!!(box()&&box().checked);const r=d.reporting;const label=ex?" (excluding flagged)":"";let cards;
if(r){const fys=r.fiscal_years.slice(-10).map(x=>({k:x.fiscal_year.replace(/^FY/,"")+(x.partial?"*":""),n:ex?x.publications_excluding_flagged:x.publications}));
const partialNote=r.fiscal_years.some(x=>x.partial)?`<p class="table-note">* Current fiscal year, to date.</p>`:"";
const max=Math.max(...fys.map(x=>x.n),1);
const bars=fys.map(x=>`<div class="mini-bar-row"><span>${esc(x.k)}</span><i style="width:${100*x.n/max}%"></i><strong>${x.n}</strong></div>`).join("");
const qs=r.quarters.map(x=>`<tr><td title="${esc(x.period_start)} to ${esc(x.period_end)}">${esc(x.fiscal_year.replace(/^FY/,""))} ${esc(x.fiscal_quarter)}${x.partial?"*":""}</td><td>${ex?x.published_excluding_flagged:x.published}</td><td>${(ex?x.added_excluding_flagged:x.added)??"—"}</td></tr>`).join("");
const note=`<p class="table-note">Q1 Jul–Sep, Q2 Oct–Dec, Q3 Jan–Mar, Q4 Apr–Jun.${r.quarters.some(x=>x.partial)?" * To date.":""}${r.quarters.some(x=>x.added===null)&&r.added_date_note?" — "+esc(r.added_date_note):""}</p>`;
const topics=r.top_topics.slice(0,10).map(x=>`<li><span>${esc(x.topic)}</span><strong>${x.publications}</strong></li>`).join("");
cards=`<article class="observatory-card"><h3>Output by fiscal year${label}</h3><div class="mini-bars fy">${bars}</div>${partialNote}</article><article class="observatory-card"><h3>Latest fiscal quarters${label}</h3><table><thead><tr><th>Quarter</th><th>Published</th><th>Added to OpenAlex</th></tr></thead><tbody>${qs}</tbody></table>${note}</article><article class="observatory-card"><h3>Leading OpenAlex topics, ${esc(r.topics_fiscal_year)}</h3><ol class="topic-rank">${topics}</ol></article>`;}
else{const annualAll=ex&&d.annual_excluding_flagged?d.annual_excluding_flagged:d.annual;const quartersAll=ex&&d.quarters_excluding_flagged?d.quarters_excluding_flagged:d.quarters;
const recent=annualAll.filter(x=>x.publication_year>=2016).slice(-10);const max=Math.max(...recent.map(x=>x.publications),1);
const annual=recent.map(x=>`<div class="mini-bar-row"><span>${x.publication_year}</span><i style="width:${100*x.publications/max}%"></i><strong>${x.publications}</strong></div>`).join("");
const qs=quartersAll.slice(-8).reverse().map(x=>`<tr><td>${esc(x.fiscal_year)}</td><td>${esc(x.fiscal_quarter)}</td><td>${esc(x.period_end)}</td><td>${x.publications}</td></tr>`).join("");
const topics=d.top_topics.slice(0,10).map(x=>`<li><span>${esc(x.topic)}</span><strong>${x.publications}</strong></li>`).join("");
cards=`<article class="observatory-card"><h3>Recent annual output${label}</h3><div class="mini-bars">${annual}</div></article><article class="observatory-card"><h3>Latest fiscal quarters${label}</h3><table><thead><tr><th>FY</th><th>Quarter</th><th>Ends</th><th>Works</th></tr></thead><tbody>${qs}</tbody></table></article><article class="observatory-card"><h3>Leading OpenAlex topics</h3><ol class="topic-rank">${topics}</ol></article>`;}
document.querySelector("#muhas-summary").innerHTML=cards;}
fetch("data/summary.json").then(r=>r.json()).then(d=>{data=d;render();if(box())box().addEventListener("change",render);}).catch(()=>{document.querySelector("#muhas-summary").textContent="Summary data are temporarily unavailable.";});})();'''
    (assets / "observatory.js").write_text(script, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--as-of", default=date.today().isoformat())
    parser.add_argument("--status", choices=["provisional", "final"], default="provisional")
    parser.add_argument("--reuse-raw", type=Path)
    parser.add_argument("--previous-raw", type=Path, help="raw JSONL.gz of the previous final snapshot")
    parser.add_argument("--previous-tag", default="", help="release tag of the previous final snapshot")
    args = parser.parse_args()
    as_of = date.fromisoformat(args.as_of)
    retrieved_at = iso_now()
    snapshot_id = f"MUHAS_OPENALEX_{START_DATE.year}_{as_of.isoformat()}_{args.status.upper()}_asof_{retrieved_at[:10]}"
    raw_path = RAW / f"{snapshot_id}.jsonl.gz"
    if args.reuse_raw:
        works = [json.loads(line) for line in gzip.open(args.reuse_raw, "rt", encoding="utf-8")]
        query_meta = {"query_filter": "reused raw snapshot", "openalex_reported_count": len(works), "pages_retrieved": None, "raw_path": str(args.reuse_raw)}
    else:
        works, query_meta = fetch_openalex(as_of, raw_path)
    work_rows, author_rows = process_works(works, retrieved_at, snapshot_id)
    annual_rows, quarter_rows, type_rows = summaries(work_rows)
    manifest = {
        "snapshot_id": snapshot_id,
        "pipeline_version": PIPELINE_VERSION,
        "snapshot_status": args.status,
        "institution_name": INSTITUTION_NAME,
        "alternate_name": INSTITUTION_ALT_NAME,
        "openalex_id": OPENALEX_ID,
        "ror": ROR,
        "period_start": START_DATE.isoformat(),
        "as_of_date": as_of.isoformat(),
        "retrieved_at": retrieved_at,
        "works_processed": len(work_rows),
        "authorship_rows": len(author_rows),
        **query_meta,
    }
    if args.previous_raw:
        changes = compare_works(load_raw(args.previous_raw), works, query_api=not args.reuse_raw)
        manifest["changes_since_previous"] = {"previous_snapshot": args.previous_tag or args.previous_raw.name, **changes}
    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    dated_csv = DOWNLOADS / f"{snapshot_id}_works.csv"
    dated_auth = DOWNLOADS / f"{snapshot_id}_authorships.csv"
    dated_db = DOWNLOADS / f"{snapshot_id}.sqlite"
    dated_xlsx = DOWNLOADS / f"{snapshot_id}.xlsx"
    write_csv(dated_csv, work_rows)
    write_csv(dated_auth, author_rows)
    write_sqlite(dated_db, work_rows, author_rows, manifest)
    write_excel(dated_xlsx, work_rows, author_rows, annual_rows, quarter_rows, type_rows, manifest)
    manifest["quarterly_report"] = write_quarterly_report(DOWNLOADS / f"{snapshot_id}_quarterly_report.xlsx", work_rows, manifest)
    for source, stable in (
        (dated_csv, DOWNLOADS / "MUHAS_Publications_LATEST.csv"),
        (dated_auth, DOWNLOADS / "MUHAS_Authorships_LATEST.csv"),
        (dated_db, DOWNLOADS / "MUHAS_Publications_LATEST.sqlite"),
        (dated_xlsx, DOWNLOADS / "MUHAS_Publications_LATEST.xlsx"),
    ):
        shutil.copy2(source, stable)
    (DATA / "snapshot_manifest.json").parent.mkdir(parents=True, exist_ok=True)
    (DATA / "snapshot_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    build_dashboard(work_rows, annual_rows, quarter_rows, type_rows, manifest)
    build_js()
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
    from build_site_seo import optimise_site
    optimise_site()

